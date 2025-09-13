"""
AiPM Phase 1 Backup Script (API Version)
----------------------------------------
Enhanced backup script that uses AiPM API endpoints to fetch real data.
Run this script manually to export all AiPM data into Excel format.

Requirements:
- pip install requests openpyxl
- Valid admin credentials for AiPM
- AiPM server running on localhost:5000
"""

import os
import requests
from openpyxl import Workbook, load_workbook
from datetime import datetime

# === Config ===
AIPM_API_BASE = os.environ.get("AIPM_API_BASE", "http://localhost:5000")
BACKUP_FILE = f"aipm_backup_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
TIMESTAMP = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

# Login credentials - load from environment variables for security
ADMIN_EMAIL = os.environ.get("AIPM_ADMIN_EMAIL")
ADMIN_PASSWORD = os.environ.get("AIPM_ADMIN_PASSWORD")

if not ADMIN_EMAIL or not ADMIN_PASSWORD:
    print("❌ ERROR: Please set AIPM_ADMIN_EMAIL and AIPM_ADMIN_PASSWORD environment variables")
    print("Example:")
    print("  export AIPM_ADMIN_EMAIL=admin@metro-construction.com")
    print("  export AIPM_ADMIN_PASSWORD=your_secure_password")
    exit(1)

def authenticate():
    """Authenticate with AiPM and get access token"""
    print("🔐 Authenticating with AiPM...")
    
    login_data = {
        "email": ADMIN_EMAIL,
        "password": ADMIN_PASSWORD
    }
    
    response = requests.post(f"{AIPM_API_BASE}/api/auth/login", json=login_data)
    
    if response.status_code != 200:
        raise Exception(f"Authentication failed: {response.status_code} - {response.text}")
    
    auth_data = response.json()
    print(f"✅ Authenticated as: {auth_data['user']['email']} ({auth_data['user']['role']})")
    return auth_data['accessToken']

def fetch_backup_data(access_token):
    """Fetch all organizational data from AiPM backup endpoint"""
    print("📊 Fetching backup data from AiPM...")
    
    headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json"
    }
    
    response = requests.get(f"{AIPM_API_BASE}/api/admin/backup", headers=headers)
    
    if response.status_code != 200:
        raise Exception(f"Backup fetch failed: {response.status_code}")
    
    if not response.text.strip():
        raise Exception("Empty response from backup endpoint")
    
    try:
        backup_data = response.json()
        print(f"✅ Fetched backup data: {backup_data['summary']}")
        return backup_data
    except Exception as e:
        raise Exception(f"Failed to parse JSON response: {e}")

def create_excel_backup(backup_data):
    """Convert API backup data to structured Excel file"""
    print(f"📝 Creating Excel backup: {BACKUP_FILE}")
    
    # Create or load workbook
    if os.path.exists(BACKUP_FILE):
        wb = load_workbook(BACKUP_FILE)
        print(f"🔄 Updating existing backup: {BACKUP_FILE}")
    else:
        wb = Workbook()
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])
        print(f"✅ Created new backup file: {BACKUP_FILE}")
    
    # Projects Tab
    if "Projects" not in wb.sheetnames:
        wb.create_sheet("Projects")
    ws_projects = wb["Projects"]
    ws_projects.delete_rows(1, ws_projects.max_row)  # Clear existing data
    
    # Projects header
    project_headers = [
        "Project ID", "Project Number", "Project Name", "Client/GC",
        "Location", "Contract Value", "Budget Value", "Overhead/Fee",
        "Status", "Created Date", "Description"
    ]
    ws_projects.append(project_headers)
    
    # Projects data
    for project in backup_data['data']['projects']:
        row = [
            project.get('id', ''),
            project.get('projectNumber', ''),
            project.get('name', ''),
            project.get('client', ''),
            project.get('address', ''),
            project.get('contractValue', ''),
            project.get('budget', ''),
            project.get('overheadFee', ''),
            project.get('status', ''),
            project.get('createdAt', ''),
            project.get('description', '')
        ]
        ws_projects.append(row)
    
    # Vendors Tab
    if "Vendors" not in wb.sheetnames:
        wb.create_sheet("Vendors")
    ws_vendors = wb["Vendors"]
    ws_vendors.delete_rows(1, ws_vendors.max_row)
    
    vendor_headers = [
        "Vendor ID", "Company", "Contact Name", "Email", "Phone", "Address", "Created Date"
    ]
    ws_vendors.append(vendor_headers)
    
    for vendor in backup_data['data']['vendors']:
        row = [
            vendor.get('id', ''),
            vendor.get('company', ''),
            vendor.get('name', ''),
            vendor.get('email', ''),
            vendor.get('phone', ''),
            vendor.get('address', ''),
            vendor.get('createdAt', '')
        ]
        ws_vendors.append(row)
    
    # Purchase Orders Tab
    if "Purchase Orders" not in wb.sheetnames:
        wb.create_sheet("Purchase Orders")
    ws_pos = wb["Purchase Orders"]
    ws_pos.delete_rows(1, ws_pos.max_row)
    
    po_headers = [
        "PO ID", "Project ID", "Vendor ID", "PO Number", "Status",
        "Total Amount", "PO Date", "Notes", "Created Date"
    ]
    ws_pos.append(po_headers)
    
    for po in backup_data['data']['purchaseOrders']:
        row = [
            po.get('id', ''),
            po.get('projectId', ''),
            po.get('vendorId', ''),
            po.get('number', ''),
            po.get('status', ''),
            po.get('totalAmount', ''),
            po.get('poDate', ''),
            po.get('notes', ''),
            po.get('createdAt', '')
        ]
        ws_pos.append(row)
    
    # Invoices Tab
    if "Invoices" not in wb.sheetnames:
        wb.create_sheet("Invoices")
    ws_invoices = wb["Invoices"]
    ws_invoices.delete_rows(1, ws_invoices.max_row)
    
    invoice_headers = [
        "Invoice ID", "Project ID", "PO ID", "Vendor Name", "Invoice Number",
        "Invoice Date", "Total Amount", "Status", "Match Status", "Created Date"
    ]
    ws_invoices.append(invoice_headers)
    
    for invoice in backup_data['data']['invoices']:
        row = [
            invoice.get('id', ''),
            invoice.get('projectId', ''),
            invoice.get('poId', ''),
            invoice.get('vendorName', ''),
            invoice.get('invoiceNumber', ''),
            invoice.get('invoiceDate', ''),
            invoice.get('totalAmount', ''),
            invoice.get('status', ''),
            invoice.get('matchStatus', ''),
            invoice.get('createdAt', '')
        ]
        ws_invoices.append(row)
    
    # Project Materials Tab
    if "Project Materials" not in wb.sheetnames:
        wb.create_sheet("Project Materials")
    ws_materials = wb["Project Materials"]
    ws_materials.delete_rows(1, ws_materials.max_row)
    
    material_headers = [
        "Material ID", "Project ID", "Category", "Model", "Description",
        "Unit", "Quantity", "Unit Price", "Cost Code", "Phase Code", "Source"
    ]
    ws_materials.append(material_headers)
    
    for material in backup_data['data']['projectMaterials']:
        row = [
            material.get('id', ''),
            material.get('projectId', ''),
            material.get('category', ''),
            material.get('model', ''),
            material.get('description', ''),
            material.get('unit', ''),
            material.get('qty', ''),
            material.get('unitPrice', ''),
            material.get('costCode', ''),
            material.get('phaseCode', ''),
            material.get('source', '')
        ]
        ws_materials.append(row)
    
    # Backup History Tab
    if "BackupHistory" not in wb.sheetnames:
        wb.create_sheet("BackupHistory")
    ws_history = wb["BackupHistory"]
    
    # Add backup record
    backup_record = [
        TIMESTAMP, 
        "API Backup executed",
        backup_data['summary']['projectCount'],
        backup_data['summary']['vendorCount'],
        backup_data['summary']['poCount'],
        backup_data['summary']['invoiceCount'],
        backup_data['summary']['materialCount']
    ]
    ws_history.append(backup_record)
    
    # Save workbook
    wb.save(BACKUP_FILE)
    print(f"📁 Backup saved as {BACKUP_FILE} at {TIMESTAMP}")
    return BACKUP_FILE

def main():
    """Main backup execution"""
    try:
        print("🚀 Starting AiPM Backup Process...")
        
        # Step 1: Authenticate
        access_token = authenticate()
        
        # Step 2: Fetch backup data
        backup_data = fetch_backup_data(access_token)
        
        # Step 3: Create Excel backup
        backup_file = create_excel_backup(backup_data)
        
        # Step 4: Summary
        print("\n" + "="*50)
        print("✅ BACKUP COMPLETED SUCCESSFULLY")
        print("="*50)
        print(f"📊 Organization ID: {backup_data['organizationId']}")
        print(f"📅 Backup Timestamp: {backup_data['timestamp']}")
        print(f"📁 Backup File: {backup_file}")
        print("\nData Summary:")
        for key, count in backup_data['summary'].items():
            print(f"  • {key}: {count}")
        print("="*50)
        
    except Exception as e:
        print(f"❌ BACKUP FAILED: {str(e)}")
        raise

if __name__ == "__main__":
    main()